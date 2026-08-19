from __future__ import annotations

import html
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
STORE = ROOT / "app/src/main/java/com/indianservers/aichemistryexplorer/LessonPageStore.kt"
OUTPUT_DIR = ROOT / "outputs/lesson-content-review"
OUTPUT_FILE = OUTPUT_DIR / "lesson_pages_review.xlsx"


CLASS_MODULES = [
    (
        "Class 6-8",
        "Class 6 Materials, Separation and Changes",
        [
            "Materials by appearance and hardness",
            "Soluble and insoluble substances",
            "Floating and sinking materials",
            "Handpicking, sieving and winnowing",
            "Sedimentation, decantation, filtration and evaporation",
            "Reversible and irreversible changes",
        ],
    ),
    (
        "Class 6-8",
        "Class 7 Acids, Metals and Everyday Reactions",
        [
            "Natural indicators and litmus",
            "Neutralisation in daily life",
            "Rusting of iron",
            "Preventing corrosion",
            "Metals and non-metals around us",
            "Reactions of metals with oxygen, water and acids",
        ],
    ),
    (
        "Class 6-8",
        "Class 8 Electric Current and Chemical Effects",
        [
            "Conductors and insulators in liquids",
            "Electrolytes and simple testers",
            "Simple electroplating for class 8",
        ],
    ),
    (
        "Class 6-8",
        "Fibres, Plastics, Fuels and Combustion",
        [
            "Natural and synthetic fibres",
            "Thermoplastics and thermosets",
            "Coal and petroleum fractions",
            "Combustion and flame",
            "Fuel efficiency and pollution",
        ],
    ),
    (
        "Class 9-10",
        "Class 9 Matter and Mixtures",
        [
            "Evaporation and diffusion",
            "Sublimation and latent heat",
            "Solutions, suspensions and colloids",
            "Tyndall effect and concentration",
            "Separating mixtures in class 9",
        ],
    ),
    (
        "Class 9-10",
        "Class 9 Atoms and Formulae",
        [
            "Laws of chemical combination",
            "Dalton atomic theory at school level",
            "Atomic mass and molecular mass",
            "Valency and writing formulae",
            "Mole concept introduction",
        ],
    ),
    (
        "Class 9-10",
        "Class 10 Chemical Substances and Carbon",
        [
            "Chemical equations and conservation",
            "Important salts in daily life",
            "Metals, metallurgy and corrosion",
            "Carbon bonding and homologous series",
            "Ethanol, ethanoic acid, soaps and detergents",
        ],
    ),
    (
        "Class 9-10",
        "Periodic Classification and Electrolysis",
        [
            "Modern periodic law",
            "Periodic trends at secondary level",
            "Chemical effects of electric current",
            "Electrolysis and electrode products",
            "Electroplating",
        ],
    ),
    (
        "Class 9-10",
        "School Practical Chemistry",
        [
            "Preparing and diluting solutions",
            "Testing gases",
            "Identifying acids and bases",
            "Precipitation observations",
            "Recording results and uncertainty",
        ],
    ),
    (
        "Class 11-12",
        "Intermediate First Year Foundations",
        [
            "Some basic concepts of chemistry",
            "Stoichiometry and limiting reagent",
            "Atomic structure for class 11",
            "Hydrogen and its compounds",
            "Environmental chemistry for school",
        ],
    ),
    (
        "Class 11-12",
        "Intermediate First Year Physical Chemistry",
        [
            "States of matter and gas laws",
            "Thermodynamics at school level",
            "Chemical and ionic equilibrium",
            "Redox reactions and oxidation number",
            "Organic chemistry basic principles",
        ],
    ),
    (
        "Class 11-12",
        "Intermediate Second Year Inorganic and Physical",
        [
            "Solutions and Raoult's Law",
            "Colligative properties",
            "d- and f-block elements",
            "Coordination compounds exam basics",
            "Electrochemistry cells and Nernst equation",
        ],
    ),
    (
        "Class 11-12",
        "Intermediate Second Year Organic and Biomolecules",
        [
            "Haloalkanes and haloarenes",
            "Alcohols, phenols and ethers",
            "Aldehydes, ketones and carboxylic acids",
            "Amines and Diazonium Salts",
            "Biomolecules, vitamins and polymers",
        ],
    ),
]


PHASE_TITLES = ["Discover", "Understand", "Observe", "Think", "Remember"]


@dataclass
class Brief:
    discoverTitle: str
    hook: str
    whyItMatters: str
    dailyLink: str
    simpleMeaning: str
    analogy: str
    keyWords: str
    memoryTrick: str
    homeExample: str
    schoolExample: str
    natureExample: str
    activity: str
    prediction: str
    reason: str
    misconception: str
    factOne: str
    factTwo: str
    factThree: str
    examAnswer: str
    takeaway: str


def slug(value: str) -> str:
    stable = {
        "Solutions and Raoult's Law": "solutions-and-raoult-law",
        "d- and f-block elements": "d-and-f-block-elements",
        "Alcohols, phenols and ethers": "alcohols-phenols-and-ethers",
        "Aldehydes, ketones and carboxylic acids": "aldehydes-ketones-and-carboxylic-acids",
        "Amines and Diazonium Salts": "amines-and-diazonium-salts",
        "Biomolecules, vitamins and polymers": "biomolecules-vitamins-and-polymers",
    }
    if value in stable:
        return stable[value]
    value = value.replace("Raoult's law", "Raoult law")
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def parse_briefs() -> dict[str, Brief]:
    source = STORE.read_text(encoding="utf-8")
    pattern = re.compile(r'"([^"]+)" -> SchoolLessonBrief\((.*?)\n    \)', re.S)
    briefs: dict[str, Brief] = {}
    for title, body in pattern.findall(source):
        fields = dict(re.findall(r"(\w+)\s*=\s*\"((?:[^\"\\]|\\.)*)\"", body, re.S))
        if len(fields) >= 20:
            briefs[title] = Brief(**{name: fields[name].replace('\\"', '"') for name in Brief.__annotations__})
    senior_pattern = re.compile(r'"([^"]+)" -> seniorBrief\((.*?)\n    \)', re.S)
    ordered_fields = [
        "discoverTitle",
        "hook",
        "whyItMatters",
        "simpleMeaning",
        "analogy",
        "keyWords",
        "memoryTrick",
        "homeExample",
        "schoolExample",
        "natureExample",
        "activity",
        "prediction",
        "reason",
        "misconception",
        "factOne",
        "factTwo",
        "factThree",
        "examAnswer",
        "takeaway",
    ]
    for title, body in senior_pattern.findall(source):
        values = [item.replace('\\"', '"') for item in re.findall(r'"((?:[^"\\]|\\.)*)"', body, re.S)]
        if len(values) >= len(ordered_fields):
            payload = dict(zip(ordered_fields, values))
            payload["dailyLink"] = values[19] if len(values) > 19 else "Before reading, notice one related substance, reaction, diagram or measurement in the examples below."
            briefs[title] = Brief(**payload)
    return briefs


def html_page(phase: str, style: str, heading: str, paragraphs: list[str], box_title: str, box_body: str) -> str:
    result = ["<section class='lesson-page'>"]
    result.append(f"<p class='phase'>Phase: {html.escape(phase)}</p>")
    result.append(f"<h2>{html.escape(heading)}</h2>")
    result.extend(f"<p>{paragraph}</p>" for paragraph in paragraphs)
    result.append(f"<div class='teacher-box'><h3>{html.escape(box_title)}</h3><p>{box_body}</p></div>")
    result.append("</section>")
    return "".join(result)


def content_for_page(lesson: str, brief: Brief, page_no: int) -> str:
    phase = PHASE_TITLES[page_no - 1]
    style = ""
    if phase == "Discover":
        return html_page(
            phase,
            style,
            brief.discoverTitle,
            [
                html.escape(brief.hook),
                f"Today we learn <b>{html.escape(lesson)}</b>. This topic matters because {html.escape(brief.whyItMatters)}.",
                "Goal: by the end, you should be able to explain it using one real example from your own life.",
            ],
            "Daily-life link",
            html.escape(brief.dailyLink),
        )
    if phase == "Understand":
        return html_page(
            phase,
            style,
            "The idea in simple words",
            [
                html.escape(brief.simpleMeaning),
                f"Easy picture: {html.escape(brief.analogy)}",
                f"Key words: <b>{html.escape(brief.keyWords)}</b>",
                f"Memory trick: <b>{html.escape(brief.memoryTrick)}</b>",
            ],
            "Teacher tip",
            "Do not memorise the sentence first. First make the picture in your mind. Then say it in your own words.",
        )
    if phase == "Observe":
        return html_page(
            phase,
            style,
            "See it in real life",
            [
                f"Home example: {html.escape(brief.homeExample)}",
                f"School example: {html.escape(brief.schoolExample)}",
                f"Nature example: {html.escape(brief.natureExample)}",
            ],
            "Mini activity",
            brief.activity,
        )
    if phase == "Think":
        return html_page(
            phase,
            style,
            "Build scientific thinking",
            [
                f"Prediction: {html.escape(brief.prediction)}",
                f"Why it happens: {html.escape(brief.reason)}",
                f"Common mistake: {html.escape(brief.misconception)}",
            ],
            "Fix the fear",
            f"For {html.escape(lesson)}, check the common mistake above first. Then compare your answer with the observation and try again.",
        )
    return html_page(
        phase,
        "Revision",
        "Quick revision",
        [
            f"Key point 1: {html.escape(brief.factOne)}",
            f"Key point 2: {html.escape(brief.factTwo)}",
            f"Key point 3: {html.escape(brief.factThree)}",
        ],
        "Exam-ready answer",
        f"{html.escape(brief.examAnswer)}<br><br>One-line takeaway: <b>{html.escape(brief.takeaway)}</b>",
    )


def plain_text(html_content: str) -> str:
    text = re.sub(r"<br\s*/?>", "\n", html_content)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", html.unescape(text)).strip()
    return re.sub(r"\s+([.,;:!?])", r"\1", text)


def audit_page(page_title: str, html_content: str) -> str:
    issues: list[str] = []
    if len(html_content) < 380:
        issues.append("Too short")
    if f"Phase: {page_title}" not in html_content:
        issues.append("Missing phase marker")
    if re.search(r"undergraduate|postgraduate", html_content, re.I):
        issues.append("Advanced wording")
    return "; ".join(issues) or "OK"


def style_sheet(ws, widths: dict[str, int]) -> None:
    header_fill = PatternFill("solid", fgColor="102A43")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7C9D9")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_workbook() -> Path:
    briefs = parse_briefs()
    rows = []
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    for level, module, lessons in CLASS_MODULES:
        for lesson in lessons:
            brief = briefs[lesson]
            lesson_id = f"{slug(module)}__{slug(lesson)}"
            for page_no, page_title in enumerate(PHASE_TITLES, start=1):
                html_content = content_for_page(lesson, brief, page_no)
                rows.append(
                    {
                        "level": level,
                        "module_title": module,
                        "lesson_title": lesson,
                        "lesson_id": lesson_id,
                        "page_no": page_no,
                        "page_title": page_title,
                        "page_content_html": html_content,
                        "page_content_text": plain_text(html_content),
                        "updated_at_export": now,
                        "audit_status": audit_page(page_title, html_content),
                    }
                )

    wb = Workbook()
    ws = wb.active
    ws.title = "lesson_pages"
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row[h] for h in headers])
    style_sheet(
        ws,
        {
            "A": 12,
            "B": 34,
            "C": 32,
            "D": 42,
            "E": 9,
            "F": 14,
            "G": 78,
            "H": 78,
            "I": 22,
            "J": 16,
        },
    )

    qa = wb.create_sheet("QA Summary")
    qa.append(["Metric", "Value"])
    lessons_count = len({row["lesson_id"] for row in rows})
    ok_pages = sum(1 for row in rows if row["audit_status"] == "OK")
    qa_rows = [
        ("Source", str(STORE.relative_to(ROOT))),
        ("Export type", "Generated SQLite lesson_pages content"),
        ("Lessons", lessons_count),
        ("Rows / pages", len(rows)),
        ("Expected pages per lesson", 5),
        ("OK pages", ok_pages),
        ("Issue pages", len(rows) - ok_pages),
        ("Generated at UTC", now),
    ]
    for item in qa_rows:
        qa.append(list(item))
    style_sheet(qa, {"A": 28, "B": 70})

    issues = wb.create_sheet("Issues")
    issues.append(["lesson_id", "lesson_title", "page_no", "page_title", "issue"])
    for row in rows:
        if row["audit_status"] != "OK":
            issues.append([row["lesson_id"], row["lesson_title"], row["page_no"], row["page_title"], row["audit_status"]])
    if issues.max_row == 1:
        issues.append(["-", "-", "-", "-", "No issues found"])
    style_sheet(issues, {"A": 42, "B": 32, "C": 10, "D": 16, "E": 38})

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    output = build_workbook()
    print(output)
